import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from datetime import datetime
from models.bill import Bill
from models.patient import Patient
from utils.logger import logger

class ReportsController:
    def __init__(self, view, parent_widget=None):
        self.view = view
        self.parent = parent_widget
        self.bind_signals()

    def bind_signals(self):
        self.view.generate_report_requested.connect(self.generate_report)
        self.view.export_excel_requested.connect(self.export_excel)
        self.view.export_pdf_requested.connect(self.export_pdf)

    def generate_report(self, report_type, start_date, end_date):
        """Fetches data and populates the corresponding table in the view."""
        logger.info(f"Generating report: {report_type} from {start_date} to {end_date}")
        try:
            if report_type == "general_revenue":
                data = Bill.get_general_revenue_report(start_date, end_date)
                self.view.populate_general_revenue(data)
            elif report_type == "doctor_revenue":
                data = Bill.get_doctor_wise_revenue_report(start_date, end_date)
                self.view.populate_doctor_revenue(data)
            elif report_type == "procedure_revenue":
                data = Bill.get_procedure_revenue_report(start_date, end_date)
                self.view.populate_procedure_revenue(data)
            elif report_type == "lab_revenue":
                data = Bill.get_lab_revenue_report(start_date, end_date)
                self.view.populate_lab_revenue(data)
            elif report_type == "patient_registrations":
                # For registrations, we fetch patient records in range
                from database.db_manager import DatabaseManager
                db = DatabaseManager()
                query = "SELECT * FROM patients WHERE date(registration_date) BETWEEN date(?) AND date(?) ORDER BY registration_date DESC;"
                with db.get_connection() as conn:
                    cursor = conn.cursor()
                    cursor.execute(query, (start_date, end_date))
                    rows = cursor.fetchall()
                    data = [dict(row) for row in rows]
                self.view.populate_registrations_report(data)
            
            from views.widgets.toast import ToastNotification
            ToastNotification.show_success("Report generated successfully.", self.parent)
        except Exception as e:
            logger.error(f"Error generating report {report_type}: {str(e)}")
            from views.widgets.toast import ToastNotification
            ToastNotification.show_error("Failed to generate report data.", self.parent)

    def export_excel(self, report_type, start_date, end_date):
        """Compiles and exports report data into a styled Excel file."""
        from PyQt6.QtWidgets import QFileDialog
        
        default_filename = f"BMC_{report_type}_report_{start_date}_to_{end_date}.xlsx"
        save_path, _ = QFileDialog.getSaveFileName(
            self.parent,
            "Export to Excel",
            default_filename,
            "Excel Files (*.xlsx)"
        )
        if not save_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Report Data"

            # Stylings
            title_font = Font(name="Segoe UI", size=14, bold=True, color="0F6CBD")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="0F6CBD", end_color="0F6CBD", fill_type="solid")
            data_font = Font(name="Segoe UI", size=10)
            align_center = Alignment(horizontal="center", vertical="center")
            align_right = Alignment(horizontal="right", vertical="center")
            border_thin = Border(
                left=Side(style='thin', color='E2E8F0'),
                right=Side(style='thin', color='E2E8F0'),
                top=Side(style='thin', color='E2E8F0'),
                bottom=Side(style='thin', color='E2E8F0')
            )

            # Retrieve report title and headers
            title_text = f"Bismel Medical Center - {report_type.replace('_', ' ').title()} Report"
            subtitle_text = f"Report Date Range: {start_date} to {end_date}"

            # Write titles
            ws.append([title_text])
            ws.cell(1, 1).font = title_font
            ws.append([subtitle_text])
            ws.cell(2, 1).font = Font(name="Segoe UI", size=11, italic=True)
            ws.append([]) # Empty row

            headers = []
            data_rows = []

            # Fetch relevant data matching table representations
            if report_type == "general_revenue":
                headers = ["Transaction Date", "Invoices Count", "Subtotal (Rs)", "Discount (Rs)", "Net Grand Total (Rs)"]
                data = Bill.get_general_revenue_report(start_date, end_date)
                for r in data:
                    data_rows.append([
                        r.get("transaction_date", ""),
                        r.get("bills_count", 0),
                        r.get("subtotal_sum", 0.0),
                        r.get("discount_sum", 0.0),
                        r.get("grand_total_sum", 0.0)
                    ])
            elif report_type == "doctor_revenue":
                headers = ["Doctor Name", "Consultations Count", "Total Revenue (Rs)"]
                data = Bill.get_doctor_wise_revenue_report(start_date, end_date)
                for r in data:
                    data_rows.append([r.get("doctor_name", ""), r.get("consultation_count", 0), r.get("total_revenue", 0.0)])
            elif report_type == "procedure_revenue":
                headers = ["Procedure Name", "Category", "Quantity Performed", "Total Revenue (Rs)"]
                data = Bill.get_procedure_revenue_report(start_date, end_date)
                for r in data:
                    data_rows.append([r.get("procedure_name", ""), r.get("category", ""), r.get("quantity_sold", 0), r.get("total_revenue", 0.0)])
            elif report_type == "lab_revenue":
                headers = ["Lab Test Name", "Quantity Performed", "Total Revenue (Rs)"]
                data = Bill.get_lab_revenue_report(start_date, end_date)
                for r in data:
                    data_rows.append([r.get("test_name", ""), r.get("quantity_sold", 0), r.get("total_revenue", 0.0)])
            elif report_type == "patient_registrations":
                headers = ["Patient ID", "Patient Name", "Gender", "Age", "Phone", "Registration Date"]
                from database.db_manager import DatabaseManager
                db = DatabaseManager()
                with db.get_connection() as conn:
                    cursor = conn.cursor()
                    cursor.execute("SELECT * FROM patients WHERE date(registration_date) BETWEEN date(?) AND date(?) ORDER BY registration_date DESC;", (start_date, end_date))
                    rows = cursor.fetchall()
                for r in rows:
                    data_rows.append([r["patient_id"], r["patient_name"], r["gender"], r["age"], r["phone"], r["registration_date"]])

            # Write headers
            ws.append(headers)
            header_row_idx = 4
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=header_row_idx, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = align_center

            # Write data
            for r_idx, row_data in enumerate(data_rows, start=header_row_idx + 1):
                ws.append(row_data)
                for c_idx in range(1, len(row_data) + 1):
                    cell = ws.cell(row=r_idx, column=c_idx)
                    cell.font = data_font
                    cell.border = border_thin
                    
                    # Align numbers nicely
                    val = row_data[c_idx - 1]
                    if isinstance(val, (int, float)):
                        cell.alignment = align_right
                        if isinstance(val, float):
                            cell.number_format = 'Rs#,##0.00'
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

            # Auto-fit columns
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(save_path)
            logger.info(f"Exported Excel report successfully to {save_path}")
            from views.widgets.toast import ToastNotification
            ToastNotification.show_success("Excel report exported successfully.", self.parent)

        except Exception as e:
            logger.error(f"Excel export failed: {str(e)}")
            from views.widgets.toast import ToastNotification
            ToastNotification.show_error("Excel export failed.", self.parent)

    def export_pdf(self, report_type, start_date, end_date):
        """Compiles and exports report details to a professional PDF document."""
        from PyQt6.QtWidgets import QFileDialog
        
        default_filename = f"BMC_{report_type}_report_{start_date}_to_{end_date}.pdf"
        save_path, _ = QFileDialog.getSaveFileName(
            self.parent,
            "Export to PDF",
            default_filename,
            "PDF Files (*.pdf)"
        )
        if not save_path:
            return

        try:
            # Report details compilation
            doc = SimpleDocTemplate(
                save_path,
                pagesize=landscape(letter) if report_type == "patient_registrations" or report_type == "general_revenue" else letter,
                rightMargin=36,
                leftMargin=36,
                topMargin=36,
                bottomMargin=36
            )

            story = []
            styles = getSampleStyleSheet()

            primary_color = colors.HexColor("#0F6CBD")
            dark_text_color = colors.HexColor("#1E293B")
            line_color = colors.HexColor("#E2E8F0")

            title_style = ParagraphStyle(
                'DocTitle',
                parent=styles['Normal'],
                fontName='Helvetica-Bold',
                fontSize=20,
                leading=24,
                textColor=primary_color
            )
            meta_style = ParagraphStyle(
                'Meta',
                parent=styles['Normal'],
                fontName='Helvetica',
                fontSize=10,
                leading=14,
                textColor=dark_text_color
            )
            table_header_style = ParagraphStyle(
                'TableHeader',
                parent=styles['Normal'],
                fontName='Helvetica-Bold',
                fontSize=9.5,
                leading=12,
                textColor=colors.white
            )
            table_cell_style = ParagraphStyle(
                'TableCell',
                parent=styles['Normal'],
                fontName='Helvetica',
                fontSize=9,
                leading=12,
                textColor=dark_text_color
            )

            # Titles
            title_text = f"Bismel Medical Center - {report_type.replace('_', ' ').title()} Report"
            story.append(Paragraph(f"<b>{title_text}</b>", title_style))
            story.append(Spacer(1, 5))
            story.append(Paragraph(f"<b>Date Range:</b> {start_date} to {end_date} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Generated On:</b> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", meta_style))
            story.append(Spacer(1, 15))

            headers = []
            data_rows = []
            col_widths = []

            # Populate data matching columns
            if report_type == "general_revenue":
                headers = ["Transaction Date", "Invoices Count", "Subtotal Sum (Rs)", "Discount Sum (Rs)", "Net Grand Total (Rs)"]
                data = Bill.get_general_revenue_report(start_date, end_date)
                for r in data:
                    data_rows.append([
                        r.get("transaction_date", ""),
                        str(r.get("bills_count", 0)),
                        f"{r.get('subtotal_sum', 0.0):,.2f}",
                        f"{r.get('discount_sum', 0.0):,.2f}",
                        f"{r.get('grand_total_sum', 0.0):,.2f}"
                    ])
                col_widths = [1.8 * 72, 1.2 * 72, 1.8 * 72, 1.8 * 72, 1.8 * 72] # landscape
            elif report_type == "doctor_revenue":
                headers = ["Doctor Name", "Consultations Count", "Total Revenue (Rs)"]
                data = Bill.get_doctor_wise_revenue_report(start_date, end_date)
                for r in data:
                    data_rows.append([r.get("doctor_name", ""), str(r.get("consultation_count", 0)), f"{r.get('total_revenue', 0.0):,.2f}"])
                col_widths = [3.5 * 72, 1.5 * 72, 2.5 * 72]
            elif report_type == "procedure_revenue":
                headers = ["Procedure Name", "Category", "Quantity Performed", "Total Revenue (Rs)"]
                data = Bill.get_procedure_revenue_report(start_date, end_date)
                for r in data:
                    data_rows.append([r.get("procedure_name", ""), r.get("category", ""), str(r.get("quantity_sold", 0)), f"{r.get('total_revenue', 0.0):,.2f}"])
                col_widths = [2.8 * 72, 1.4 * 72, 1.3 * 72, 2.0 * 72]
            elif report_type == "lab_revenue":
                headers = ["Lab Test Name", "Quantity Performed", "Total Revenue (Rs)"]
                data = Bill.get_lab_revenue_report(start_date, end_date)
                for r in data:
                    data_rows.append([r.get("test_name", ""), str(r.get("quantity_sold", 0)), f"{r.get('total_revenue', 0.0):,.2f}"])
                col_widths = [3.5 * 72, 1.5 * 72, 2.5 * 72]
            elif report_type == "patient_registrations":
                headers = ["Patient ID", "Patient Name", "Gender", "Age", "Phone", "Registration Date"]
                from database.db_manager import DatabaseManager
                db = DatabaseManager()
                with db.get_connection() as conn:
                    cursor = conn.cursor()
                    cursor.execute("SELECT * FROM patients WHERE date(registration_date) BETWEEN date(?) AND date(?) ORDER BY registration_date DESC;", (start_date, end_date))
                    rows = cursor.fetchall()
                for r in rows:
                    data_rows.append([r["patient_id"], r["patient_name"], r["gender"], str(r["age"]), r["phone"], r["registration_date"]])
                col_widths = [1.2 * 72, 2.2 * 72, 0.8 * 72, 0.6 * 72, 1.5 * 72, 2.2 * 72]

            # Build elements table
            table_headers = [Paragraph(f"<b>{h}</b>", table_header_style) for h in headers]
            pdf_data = [table_headers]

            for row in data_rows:
                pdf_data.append([Paragraph(cell, table_cell_style) for cell in row])

            table_el = Table(pdf_data, colWidths=col_widths)
            
            # Apply beautiful report styling
            t_style = [
                ('BACKGROUND', (0,0), (-1,0), primary_color),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('TOPPADDING', (0,0), (-1,-1), 5),
                ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                ('GRID', (0,0), (-1,-1), 0.5, line_color),
            ]
            
            # Alternating backgrounds
            for idx in range(1, len(pdf_data)):
                if idx % 2 == 0:
                    t_style.append(('BACKGROUND', (0, idx), (-1, idx), colors.HexColor("#F8FAFC")))
                    
            table_el.setStyle(TableStyle(t_style))
            story.append(table_el)
            
            doc.build(story)
            logger.info(f"Exported PDF report successfully to {save_path}")
            from views.widgets.toast import ToastNotification
            ToastNotification.show_success("PDF report exported successfully.", self.parent)

        except Exception as e:
            logger.error(f"PDF report export failed: {str(e)}")
            from views.widgets.toast import ToastNotification
            ToastNotification.show_error("PDF report export failed.", self.parent)
